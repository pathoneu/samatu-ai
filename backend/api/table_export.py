"""GET /api/projects/{project_id}/lenteles/{kind}.xlsx — „Eksportuoti lentelę": a table the „lentelės" stage read
(Pathon v3 W3-A, `_meta.lenteles.ziniarasciai[kind]`) sent to Excel as the drawing writes it (Pathon v3 W5-D, D10).

One sheet: the table's title, its header (the column words the drawing prints), every row's cells as read, the
table's own totals, and where it stands (sheet number) — the person's check copy, not an estimate. Owner check
first, read-only, €0; every cell goes through the spreadsheet formula guard (extractor/spreadsheet_safety.py).
"""
from __future__ import annotations

import io
import re

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import StreamingResponse

from api.free_reread import _project_for
from extractor.spreadsheet_safety import sanitize_spreadsheet_cell as _safe

router = APIRouter(prefix="/api/projects", tags=["table-export"])

KIND_RE = re.compile(r"^[a-z_]{2,32}$")
KIND_LT = {"durys": "Durų žiniaraštis", "langai": "Langų žiniaraštis", "apdaila": "Apdailos lentelė",
           "plyteles": "Plytelių lentelė", "eksplikacija": "Patalpų eksplikacija"}
NO_TABLE = "Tokios lentelės šiame projekte neperskaityta — paleiskite nemokamą perskaitymą iš naujo."


def table_rows(meta: dict, kind: str) -> tuple[str, list[str], list[list], list[dict], list[int]] | None:
    """(title, header, rows of cells, totals, sheets) of one read schedule, or None. Pure."""
    lt = meta.get("lenteles") if isinstance(meta.get("lenteles"), dict) else {}
    sched = (lt.get("ziniarasciai") or {}).get(kind)
    if not isinstance(sched, dict) or not sched.get("rows"):
        return None
    header: list[str] = []
    ids = set(sched.get("table_ids") or [])
    for tables in (lt.get("lenteles") or {}).values():
        for t in tables or []:
            if t.get("table_id") in ids and t.get("header"):
                header = [str(h or "") for h in t["header"]]
                break
        if header:
            break
    if not header and isinstance(sched.get("columns"), dict):
        header = [str(v) for v in sched["columns"].values()]
    rows = [[c for c in (r.get("cells") or [])] for r in sched["rows"] if isinstance(r, dict)]
    totals = [t for t in sched.get("totals") or [] if isinstance(t, dict)]
    pages = sorted({int(p) for p in (sched.get("pages") or [sched.get("page")]) if p})
    title = str(sched.get("title") or KIND_LT.get(kind) or kind)
    return title, header, rows, totals, pages


@router.get("/{project_id}/lenteles/{kind}.xlsx")
def api_table_export(request: Request, project_id: str, kind: str):
    project = _project_for(request, project_id)             # owner first, before anything about the request
    if not KIND_RE.match(kind or ""):
        raise HTTPException(status_code=400, detail="Netinkamas lentelės pavadinimas")
    meta = project.get("_meta") if isinstance(project.get("_meta"), dict) else {}
    got = table_rows(meta, kind)
    if got is None:
        raise HTTPException(status_code=404, detail=NO_TABLE)
    title, header, rows, totals, pages = got
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = (KIND_LT.get(kind) or "Lentelė")[:31]
    ws.append([_safe(title)])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([_safe(f"Perskaityta iš brėžinio — {'lapas' if len(pages) == 1 else 'lapai'} {', '.join(map(str, pages))}"
                     if pages else "Perskaityta iš brėžinio")])
    ws.append([])
    if header:
        ws.append([_safe(h) for h in header])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for r in rows:
        ws.append([_safe(c) for c in r])
    for t in totals:
        ws.append([_safe(str(t.get("label") or "Iš viso")), _safe(t.get("value_raw") or t.get("value"))])
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="lentele-{kind}-{project_id}.xlsx"'})
