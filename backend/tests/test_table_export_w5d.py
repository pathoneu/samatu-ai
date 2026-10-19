"""„Eksportuoti lentelę" (Pathon v3 W5-D, D10): a read schedule sent to Excel as the drawing writes it."""
from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app  # noqa: E402

import auth  # noqa: E402
import db  # noqa: E402
from api.table_export import table_rows  # noqa: E402

LENTELES = {
    "lenteles": {"1": [{"table_id": "p1-t1", "page": 1, "header": ["Nr.", "PATALPA", "m2"]}]},
    "ziniarasciai": {"eksplikacija": {
        "kind": "eksplikacija", "title": "PATALPŲ EKSPLIKACIJA", "page": 1, "pages": [1], "table_ids": ["p1-t1"],
        "columns": {"nr": "Nr.", "pavadinimas": "PATALPA", "plotas": "m2"},
        "rows": [{"cells": ["01", "Tambūras", "4.48"]}, {"cells": ["02", "=SUM(A1)", "6.90"]}],
        "totals": [{"label": "BENDRAS PLOTAS", "value": 11.38, "value_raw": "11.38"}]}},
}


def test_the_table_as_read_title_header_rows_totals():
    title, header, rows, totals, pages = table_rows({"lenteles": LENTELES}, "eksplikacija")
    assert title == "PATALPŲ EKSPLIKACIJA" and header == ["Nr.", "PATALPA", "m2"] and pages == [1]
    assert rows[0] == ["01", "Tambūras", "4.48"] and totals[0]["label"] == "BENDRAS PLOTAS"
    assert table_rows({"lenteles": LENTELES}, "durys") is None and table_rows({}, "eksplikacija") is None


def _client(tmp_path, monkeypatch, blob):
    monkeypatch.setattr(db, "_DB_PATH", tmp_path / "t.db")
    monkeypatch.setattr(auth, "AUTH_DISABLED", True)
    db.init_schema()
    db._migrate_schema()
    db.upsert_project("p1", blob)
    return TestClient(app)


def test_the_route_sends_an_xlsx_with_formulas_defused(tmp_path, monkeypatch):
    c = _client(tmp_path, monkeypatch, {"name": "P", "patalpos": [], "_meta": {"lenteles": LENTELES}})
    r = c.get("/api/projects/p1/lenteles/eksplikacija.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[0][0] == "PATALPŲ EKSPLIKACIJA"
    assert ["Nr.", "PATALPA", "m2"] in values
    assert ["01", "Tambūras", "4.48"] in values
    assert not any(isinstance(v, str) and v.startswith("=") for row in values for v in row)   # the formula guard
    assert values[-1][:2] == ["BENDRAS PLOTAS", "11.38"]


def test_the_route_refuses_what_it_has_not_got(tmp_path, monkeypatch):
    c = _client(tmp_path, monkeypatch, {"name": "P", "patalpos": []})
    assert c.get("/api/projects/p1/lenteles/eksplikacija.xlsx").status_code == 404
    assert c.get("/api/projects/p1/lenteles/Bad-Kind.xlsx").status_code == 400
    assert c.get("/api/projects/nope/lenteles/eksplikacija.xlsx").status_code == 404
