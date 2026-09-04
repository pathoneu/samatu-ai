"""P2 (MED) — anonymous / cross-tenant XLSX export via ?fixture= param.

2026-07-04 pentest, pass 3, finding P2. `api/export_excel._load_project` returned
from the ``?fixture=`` branch BEFORE `require_user` ran, so
`POST /api/projects/{Bid}/export-excel?fixture=objektas_a` needed no auth and the
endpoint then called `_load_calibration({Bid})` → read
`data/calibrations/{Bid}.json` keyed by the attacker-supplied URL project_id. The
"Patikslinta Sąmata" sheet echoed B's secret calibration ``user_value`` totals
(their KNOWN category quantities) straight into a downloadable XLSX.

The fix runs `require_owned_project(request, project_id)` BEFORE the fixture
branch, so `_load_calibration` is only ever reached for a caller who owns the URL
project_id.

These tests REPLAY the PoC. Import main BEFORE auth so load_dotenv() runs first;
force real auth so the dev bypass doesn't mask the gate.
"""
from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app  # noqa: E402  (must precede `import auth`)

import auth  # noqa: E402
import db  # noqa: E402
import api.export_excel as export_mod  # noqa: E402

_PW = "passw0rd123"
# A distinctive secret value that would leak into the XLSX if the gate failed.
_SECRET_USER_VALUE = 133742.75


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "_DB_PATH", tmp_path / "p2_export.db")
    monkeypatch.setattr(auth, "AUTH_DISABLED", False)
    calib_root = tmp_path / "calibrations"
    calib_root.mkdir()
    monkeypatch.setattr(export_mod, "CALIB_ROOT", calib_root)
    db.init_schema()
    db._migrate_schema()
    owner = db.create_user("owner@x.lt", auth.hash_password(_PW), name="Owner")
    db.create_user("attacker@x.lt", auth.hash_password(_PW), name="Att")
    db.upsert_project("proj-B", {"name": "B", "tipas": "interjeras"},
                      user_id=owner["id"])
    # Seed B's PRIVATE calibration with a secret override total.
    (calib_root / "proj-B.json").write_text(json.dumps({
        "updated_at": 1234567890,
        "factors": {
            "led_nisos_m": {"applied": True, "factor": 1.5,
                            "user_value": _SECRET_USER_VALUE},
        },
    }), encoding="utf-8")
    return calib_root


def _login(email):
    c = TestClient(app)
    r = c.post("/api/auth/login", json={"email": email, "password": _PW})
    assert r.status_code == 200, r.text
    return c


def _xlsx_contains(content: bytes, needle: float) -> bool:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, (int, float)) and abs(cell - needle) < 1e-6:
                    return True
    return False


def test_p2_anon_with_fixture_rejected(env):
    """Anon POST with ?fixture= (the original PoC) → 401, no XLSX, no leak."""
    r = TestClient(app).post("/api/projects/proj-B/export-excel?fixture=objektas_a")
    assert r.status_code == 401, r.text
    # No spreadsheet body at all.
    assert "spreadsheetml" not in r.headers.get("content-type", "")
    assert str(int(_SECRET_USER_VALUE)) not in r.text


def test_p2_cross_tenant_export_with_fixture_blocked(env):
    """A logged in, targets B's project WITH ?fixture= → 404, no XLSX, no leak."""
    att = _login("attacker@x.lt")
    r = att.post("/api/projects/proj-B/export-excel?fixture=objektas_a")
    assert r.status_code == 404, r.text
    assert "spreadsheetml" not in r.headers.get("content-type", "")
    assert str(int(_SECRET_USER_VALUE)) not in r.text


def test_p2_owner_export_still_works_and_carries_own_calibration(env):
    """The owner → 200 XLSX; the Patikslinta sheet legitimately shows B's total."""
    owner = _login("owner@x.lt")
    r = owner.post("/api/projects/proj-B/export-excel?fixture=objektas_a")
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    # Calibration present → second sheet exists and carries the owner's own value.
    assert "Patikslinta Sąmata" in wb.sheetnames
    assert _xlsx_contains(r.content, _SECRET_USER_VALUE), \
        "owner's own calibration user_value should appear in their XLSX"
