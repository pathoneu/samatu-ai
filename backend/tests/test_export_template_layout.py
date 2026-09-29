"""`/export/excel?layout=template` prints HIS layout, not ours wearing his labels.

Week 2, 2026-09-04. The default export exists to tidy OUR output: it migrates
legacy section names onto the canonical pair and re-sorts sections into the
product's own order. Applied to the contractor's template that is exactly
backwards — his headings get renamed and his order rewritten — so the layout
branch turns both off and prints the rows as given.

The second property here is the one the council insisted on: a cell he has not
filled must stay EMPTY. A `=ROUND(D*E,2)` over a blank quantity prints
„0,00 €", which reads as „this work costs nothing" rather than „įrašykite" —
and this workbook goes to his client.

Every section and label below is invented.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "backend"))

import main  # noqa: E402

_PW = "passw0rd123"

# His order, and a heading whose words the canonical migration would rewrite.
HIS_ROWS = [
    {"nr": 1, "skyrius": "Omega išlaidos", "aprasymas": "Šiukšlių išvežimas",
     "mato_vnt": "m3", "kiekis": 6.0, "kaina": 30.0},
    {"nr": 2, "skyrius": "Omega išlaidos", "aprasymas": "Išgalvota nuoma",
     "mato_vnt": "d.d.", "kiekis": None, "kaina": None},
    {"nr": 3, "skyrius": "Sienų lubų glaistymas dažymas",
     "aprasymas": "Sienų, lubų glaistymas/dažymas", "mato_vnt": "m2",
     "kiekis": None, "kaina": None},
    {"nr": 4, "skyrius": "Alfa darbai", "aprasymas": "Angokraščių apdirbimas GKP",
     "mato_vnt": "m", "kiekis": 40.0, "kaina": 5.0},
]


@pytest.fixture()
def client(tmp_path, monkeypatch):
    import auth
    import db
    monkeypatch.setattr(db, "_DB_PATH", tmp_path / "export-template.db")
    monkeypatch.setattr(auth, "AUTH_DISABLED", False)
    db.init_schema()
    db._migrate_schema()
    db.create_user("owner@x.lt", auth.hash_password(_PW), name="Owner")
    c = TestClient(main.app)
    assert c.post("/api/auth/login",
                  json={"email": "owner@x.lt", "password": _PW}).status_code == 200
    return c


def _sheet(response):
    wb = load_workbook(io.BytesIO(response.content))
    return wb["Sąmata"]


def _rows(ws):
    """(A, B, D, F) per written row, headers and section bars included."""
    return [(ws.cell(r, 1).value, ws.cell(r, 2).value,
             ws.cell(r, 4).value, ws.cell(r, 6).value)
            for r in range(4, ws.max_row + 1)]


def _payload():
    return {"projektas": "Bandomasis", "eilutes": [dict(r) for r in HIS_ROWS]}


# ── The layout itself ────────────────────────────────────────────────────────

def _headers(ws):
    """Section bars: a value in A, nothing anywhere else. („Skyriaus suma:"
    also sits alone in A but carries the subtotal in F.)"""
    return [a for a, b, d, f in _rows(ws) if a and b is None and d is None
            and f is None]


def test_his_sections_keep_their_names_and_their_order(client):
    ws = _sheet(client.post("/export/excel?layout=template", json=_payload()))
    assert _headers(ws) == ["Omega išlaidos", "Sienų lubų glaistymas dažymas",
                            "Alfa darbai"]


def test_without_the_layout_flag_the_default_export_is_unchanged(client):
    """The control. Same payload, no flag: our canonical order and the paint
    section split in two — which is right for OUR output and wrong for his."""
    headers = _headers(_sheet(client.post("/export/excel", json=_payload())))
    assert headers != ["Omega išlaidos", "Sienų lubų glaistymas dažymas",
                       "Alfa darbai"]
    assert "Sienų lubų glaistymas dažymas" not in headers


def test_his_labels_are_printed_verbatim(client):
    ws = _sheet(client.post("/export/excel?layout=template", json=_payload()))
    labels = [b for a, b, d, f in _rows(ws) if b]
    assert "Angokraščių apdirbimas GKP" in labels
    assert "Išgalvota nuoma" in labels


def test_rows_are_still_renumbered_one_to_n(client):
    ws = _sheet(client.post("/export/excel?layout=template", json=_payload()))
    nrs = [a for a, b, d, f in _rows(ws) if b]
    assert nrs == ["1", "2", "3", "4"]


# ── Blank cells ──────────────────────────────────────────────────────────────

def test_a_cell_he_has_not_filled_prints_empty_not_zero(client):
    ws = _sheet(client.post("/export/excel?layout=template", json=_payload()))
    by_label = {b: (d, f) for a, b, d, f in _rows(ws) if b}
    kiekis, suma = by_label["Išgalvota nuoma"]
    assert kiekis is None
    assert suma is None, "a formula over an empty Kiekis prints 0,00 €"
    kiekis, suma = by_label["Sienų, lubų glaistymas/dažymas"]
    assert kiekis is None and suma is None


def test_a_filled_row_still_carries_its_formula(client):
    ws = _sheet(client.post("/export/excel?layout=template", json=_payload()))
    by_label = {b: (d, f) for a, b, d, f in _rows(ws) if b}
    kiekis, suma = by_label["Šiukšlių išvežimas"]
    assert kiekis == 6.0
    assert isinstance(suma, str) and suma.startswith("=ROUND(")


def test_the_default_layout_keeps_writing_its_formula_everywhere(client):
    """The blank-cell rule is scoped to his layout — nothing else moves."""
    ws = _sheet(client.post("/export/excel", json=_payload()))
    sumas = [f for a, b, d, f in _rows(ws) if b]
    assert all(isinstance(s, str) and s.startswith("=ROUND(") for s in sumas)


# ── A merged line without one unit price (audit 2026-09-29) ─────────────────
# Several of our rows fed one of HIS lines at different prices (or only some of
# them priced). The render sends the line's amount with no unit price; a
# `=ROUND(D*E,2)` over the blank Kaina would print 0,00 € and drop the money
# from the total, so the amount is written as a value and Kaina stays blank.

MIXED_ROWS = [
    {"nr": 1, "skyrius": "Alfa darbai", "aprasymas": "Išgalvotų šviestuvų montavimas",
     "mato_vnt": "vnt", "kiekis": 35.0, "kaina": None, "suma": 825.0},
    {"nr": 2, "skyrius": "Alfa darbai", "aprasymas": "Angokraščių apdirbimas GKP",
     "mato_vnt": "m", "kiekis": 40.0, "kaina": 5.0},
    {"nr": 3, "skyrius": "Omega išlaidos", "aprasymas": "Išgalvota nuoma",
     "mato_vnt": "d.d.", "kiekis": None, "kaina": None},
]


def _mixed_sheet(client):
    return _sheet(client.post("/export/excel?layout=template",
                              json={"projektas": "Bandomasis",
                                    "eilutes": [dict(r) for r in MIXED_ROWS]}))


def _row_of(ws, label):
    return next(r for r in range(4, ws.max_row + 1) if ws.cell(r, 2).value == label)


def test_a_line_without_one_unit_price_prints_its_amount_and_a_blank_price(client):
    ws = _mixed_sheet(client)
    r = _row_of(ws, "Išgalvotų šviestuvų montavimas")
    assert ws.cell(r, 4).value == 35.0
    assert ws.cell(r, 5).value is None, "a blank Kaina must stay blank, never 0,00"
    assert ws.cell(r, 6).value == 825.0
    # the neighbour keeps its live formula
    r2 = _row_of(ws, "Angokraščių apdirbimas GKP")
    assert ws.cell(r2, 6).value == f"=ROUND(D{r2}*E{r2},2)"
    # and a cell he never filled is still empty
    r3 = _row_of(ws, "Išgalvota nuoma")
    assert ws.cell(r3, 5).value is None and ws.cell(r3, 6).value is None


def test_the_template_total_is_the_sum_of_the_line_amounts(client):
    """Evaluate the sheet by hand: every line's Suma (value or its ROUND
    formula) inside the „Suma be PVM" SUBTOTAL range must add up to the money
    on the lines — 825 + 40 × 5."""
    import re as _re
    ws = _mixed_sheet(client)
    tot_r = next(r for r in range(4, ws.max_row + 1)
                 if ws.cell(r, 1).value == "Suma be PVM:")
    m = _re.fullmatch(r"=SUBTOTAL\(9,F(\d+):F(\d+)\)", ws.cell(tot_r, 6).value)
    assert m, ws.cell(tot_r, 6).value
    total = 0.0
    for r in range(int(m.group(1)), int(m.group(2)) + 1):
        v = ws.cell(r, 6).value
        if isinstance(v, (int, float)):
            total += v
        elif isinstance(v, str) and v.startswith("=ROUND(D"):
            total += round((ws.cell(r, 4).value or 0) * (ws.cell(r, 5).value or 0), 2)
        # SUBTOTAL section rows are skipped by SUBTOTAL(9, …) itself
    assert total == pytest.approx(825.0 + 200.0)
