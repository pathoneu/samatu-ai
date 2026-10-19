"""Reverse-VAT (atvirkštinis PVM) export path (2026-06-19 hunt).

Renovation projects must be billed reverse-VAT: the contractor does NOT add 21%
output VAT (the užsakovas self-accounts, PVMĮ 96 str.). The main /generate-pdf and
/generate-excel exports previously always added 21% — an illegal VAT line and a
21%-inflated total on a legally-sensitive document.
"""
import io
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))
import pdfplumber
from fastapi.testclient import TestClient
from main import app

client = TestClient(app)


def _payload(archetipas=None, vat_mode=None):
    p = {
        "projektas": "Testas",
        "tipas": "interjeras",
        "eilutes": [{"nr": "1", "aprasymas": "Sienų dažymas", "mato_vnt": "m²",
                     "kiekis": 10, "kaina": 5, "suma": 50, "skyrius": "Apdaila"}],
    }
    if archetipas:
        p["archetipas"] = archetipas
    if vat_mode:
        p["vat_mode"] = vat_mode
    return p


def _pdf_text(content):
    with pdfplumber.open(io.BytesIO(content)) as d:
        return "\n".join((pg.extract_text() or "") for pg in d.pages)


def test_pdf_standard_project_charges_21_pct_pvm():
    # (the PDF totals table splits "VISO su PVM:" across columns, so assert on the
    # discriminating fragments rather than the full label.)
    r = client.post("/generate-pdf", json=_payload("standart"))
    assert r.status_code == 200, r.text
    txt = _pdf_text(r.content)
    assert "Atvirkštinis PVM" not in txt       # NOT reverse-VAT
    assert "VISO su" in txt                     # standard "VISO su PVM" total label


def test_pdf_renovacija_is_reverse_vat_no_pvm_line():
    r = client.post("/generate-pdf", json=_payload("renovacija"))
    assert r.status_code == 200, r.text
    txt = _pdf_text(r.content)
    assert "Atvirkštinis PVM" in txt           # reverse-VAT footnote present
    assert "VISO su" not in txt                 # NOT the standard "VISO su PVM" label


def test_pdf_explicit_reverse_vat_mode_also_works():
    r = client.post("/generate-pdf", json=_payload(vat_mode="reverse_vat"))
    assert r.status_code == 200, r.text
    txt = _pdf_text(r.content)
    assert "Atvirkštinis PVM" in txt


def test_excel_renovacija_generates_without_pvm_row():
    # Excel uses live formulas; just assert it builds and the reverse-VAT branch
    # doesn't error (the totals branch is exercised).
    r = client.post("/generate-excel", json=_payload("renovacija"))
    assert r.status_code == 200, r.text
    assert r.content[:2] == b"PK"  # valid xlsx (zip)


# ── HIS choice wins over the drawing label (audit 2026-09-29) ────────────────
# `archetipas` is read off the drawing, not chosen by the contractor. An explicit
# `vat_mode` must win in both directions; with none, today's behaviour stays.

from extractor.vat_mode import is_reverse_vat  # noqa: E402


def test_helper_explicit_choice_wins_and_absent_keeps_the_label_rule():
    assert is_reverse_vat({"vat_mode": "standard", "archetipas": "renovacija"}) is False
    assert is_reverse_vat({"vat_mode": "reverse_vat", "archetipas": "standart"}) is True
    assert is_reverse_vat({"vat_mode": "reverse_vat"}) is True
    # absent → exactly as before
    assert is_reverse_vat({"archetipas": "renovacija"}) is True
    assert is_reverse_vat({"archetipas": "premium"}) is False
    assert is_reverse_vat({}) is False
    assert is_reverse_vat(None) is False


def test_pdf_explicit_standard_on_a_renovacija_label_prints_21_pct():
    r = client.post("/generate-pdf", json=_payload("renovacija", vat_mode="standard"))
    assert r.status_code == 200, r.text
    txt = _pdf_text(r.content)
    assert "Atvirkštinis PVM" not in txt
    assert "VISO su" in txt
    assert "PVM 21%" in txt


def test_pdf_explicit_reverse_on_an_ordinary_project_prints_no_pvm_line():
    r = client.post("/generate-pdf", json=_payload("standart", vat_mode="reverse_vat"))
    assert r.status_code == 200, r.text
    txt = _pdf_text(r.content)
    assert "Atvirkštinis PVM" in txt
    assert "VISO su" not in txt
    assert "PVM 21%" not in txt
    # his choice is not a renovation claim
    assert "(renovacija)" not in txt


def test_pdf_absent_mode_keeps_todays_result_for_both_labels():
    assert "VISO su" in _pdf_text(client.post("/generate-pdf", json=_payload("premium")).content)
    assert "VISO su" not in _pdf_text(client.post("/generate-pdf", json=_payload("renovacija")).content)


def _excel_labels(payload):
    from openpyxl import load_workbook
    r = client.post("/generate-excel", json=payload)
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content))["Sąmata"]
    return [ws.cell(row, 1).value for row in range(1, ws.max_row + 1)]


def test_excel_explicit_standard_on_a_renovacija_label_prints_21_pct():
    labels = _excel_labels(_payload("renovacija", vat_mode="standard"))
    assert "PVM 21%:" in labels
    assert "IŠ VISO (be PVM):" not in labels


def test_excel_explicit_reverse_on_an_ordinary_project_prints_no_pvm_line():
    labels = _excel_labels(_payload("standart", vat_mode="reverse_vat"))
    assert "PVM 21%:" not in labels
    assert "IŠ VISO (be PVM):" in labels


def test_excel_absent_mode_keeps_todays_result_for_both_labels():
    assert "PVM 21%:" in _excel_labels(_payload("premium"))
    assert "PVM 21%:" not in _excel_labels(_payload("renovacija"))


# ── His rate reaches both exports (v3 W7-A M4, 2026-10-18) ──────────────────
# Settings › PVM tarifas offered 21 / 9 / 0 % but both exports printed 21 %.

import pytest  # noqa: E402

from extractor.vat_mode import vat_rate_percent  # noqa: E402


def _rated(rate, **kw):
    p = _payload(kw.pop("archetipas", None), vat_mode=kw.pop("vat_mode", None))
    if rate is not ...:
        p["vat_rate"] = rate
    return p


@pytest.mark.parametrize("rate, expected_pvm", [(21, "10,50"), (9, "4,50"), (0, "0,00")])
def test_pdf_prints_his_rate_and_the_matching_total(rate, expected_pvm):
    r = client.post("/generate-pdf", json=_rated(rate))
    assert r.status_code == 200, r.text
    txt = _pdf_text(r.content)
    assert f"PVM {rate}%" in txt
    assert expected_pvm in txt
    total = {21: "60,50", 9: "54,50", 0: "50,00"}[rate]
    assert total in txt


@pytest.mark.parametrize("rate", [21, 9, 0])
def test_excel_prints_his_rate_with_the_matching_formula(rate):
    from openpyxl import load_workbook
    r = client.post("/generate-excel", json=_rated(rate))
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content))["Sąmata"]
    rows = {ws.cell(row, 1).value: ws.cell(row, 6).value for row in range(1, ws.max_row + 1)}
    assert f"PVM {rate}%:" in rows
    assert f"*{rate / 100:.2f},2)" in rows[f"PVM {rate}%:"]
    assert "PVM 21%:" not in rows or rate == 21


def test_reverse_vat_wording_unchanged_whatever_the_rate():
    r = client.post("/generate-pdf", json=_rated(9, vat_mode="reverse_vat"))
    txt = _pdf_text(r.content)
    assert "Atvirkštinis PVM" in txt and "PVM 9%" not in txt
    assert "IŠ VISO (be PVM):" in _excel_labels(_rated(9, vat_mode="reverse_vat"))


def test_rate_helper_reads_his_value_and_never_invents_one():
    assert vat_rate_percent({"vat_rate": 9}) == 9
    assert vat_rate_percent({"vat_rate": "0"}) == 0
    assert vat_rate_percent({"vat_rate": "9 %"}) == 9
    assert vat_rate_percent({"vat_rate": 15}) == 21        # not a rate he can pick
    assert vat_rate_percent({"vat_rate": True}) == 21
    assert vat_rate_percent({}) == 21                      # old clients: unchanged
    assert vat_rate_percent({"vat_rate": 9, "vat_mode": "reverse_vat"}) == 0


def test_chat_price_summary_follows_the_same_rule():
    """The assistant's „kiek su PVM?" must agree with the top bar and the files."""
    from main import _build_chat_context
    rows = [{"aprasymas": "Sienų dažymas", "kiekis": 10, "kaina": 5}]
    rev = _build_chat_context({"darbu_sarasas": rows, "vat_mode": "reverse_vat"})
    assert rev["kainu_suvestine"]["pvm_tarifas_proc"] == 0
    std = _build_chat_context({"darbu_sarasas": rows, "archetipas": "renovacija",
                               "vat_mode": "standard"})
    assert std["kainu_suvestine"]["pvm_tarifas_proc"] == 21
    plain = _build_chat_context({"darbu_sarasas": rows})
    assert plain["kainu_suvestine"]["pvm_tarifas_proc"] == 21
