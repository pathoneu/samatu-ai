"""F6 (LOW) — CSV / Excel formula (CSV) injection on third-party-facing exports.

2026-07-04 pentest, finding F6. The F2/F3 CSV writer and the Excel exporter wrote
user/AI strings (row labels, aprašymas, notes, room names) with no neutralisation
of a leading = + - @ (or tab/CR). F2/F3 documents are shared with third parties,
so a crafted label like =HYPERLINK("http://x","c") would evaluate as a live
formula when they open the file.

The fix adds a shared sanitize_spreadsheet_cell() helper, applied on the CSV and
openpyxl text cells. These tests REPLAY the PoC and assert the risky value is
neutralised with a leading single quote.
"""
from __future__ import annotations

import csv
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from extractor.spreadsheet_safety import sanitize_spreadsheet_cell  # noqa: E402
from extractor.f2_export import build_f2_export, export_to_csv  # noqa: E402

_ATTACK = '=HYPERLINK("http://evil.example","click me")'


def test_helper_neutralises_all_formula_triggers():
    for lead in ("=", "+", "-", "@", "\t", "\r", "\n"):
        assert sanitize_spreadsheet_cell(lead + "x") == "'" + lead + "x"
    # Benign strings + numeric cells untouched.
    assert sanitize_spreadsheet_cell("Sienų dažymas") == "Sienų dažymas"
    assert sanitize_spreadsheet_cell(230) == 230
    assert sanitize_spreadsheet_cell(14.0) == 14.0
    assert sanitize_spreadsheet_cell(None) is None


def test_f6_csv_export_neutralises_hyperlink_label():
    lines = [{
        "code": "poz-001",
        "label_lt": _ATTACK,           # attacker-controlled label
        "unit": "m²",
        "quantity": 10,
        "unit_price_eur_be_pvm": 5.0,
        "confidence": "green",
        "source": "user",
        "notes": "=1+1",               # attacker-controlled note
    }]
    export = build_f2_export("p1", lines, [], vat_mode="be_pvm")
    csv_text = export_to_csv(export)

    # Parse back out and find the data row.
    rows = list(csv.reader(io.StringIO(csv_text)))
    data_row = next(r for r in rows if r and r[0] == "1")
    label_cell = data_row[2]   # "Pavadinimas" column
    notes_cell = data_row[11]  # "Pastabos" column

    assert label_cell.startswith("'="), f"label not neutralised: {label_cell!r}"
    assert label_cell == "'" + _ATTACK
    assert notes_cell == "'=1+1"
    # The raw formula must NOT appear unquoted anywhere.
    assert "\n=HYPERLINK" not in csv_text
    assert ",=HYPERLINK" not in csv_text


def test_f6_excel_export_neutralises_project_name_and_room():
    from openpyxl import load_workbook
    import api.export_excel as ex

    parsed = {
        "pavadinimas": _ATTACK,           # attacker-controlled project name
        "santrauka": {},
        "patalpos": [
            {"nr": "=cmd", "pavadinimas": "@SUM(A1:A9)", "plotas_m2": 12.5,
             "perimetras_m": 14.0, "dazymo_plotas_m2": 30.0},
        ],
    }
    from openpyxl import Workbook
    wb = Workbook()
    ex._write_kz_sheet(wb.active, parsed, calibration=None)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    ws = load_workbook(buf).active
    assert ws["A1"].value == "'" + _ATTACK, ws["A1"].value
    # Find the room row (nr cell neutralised, name cell neutralised, area numeric).
    neutralised_nr = [c for col in ws.iter_cols() for c in col if c.value == "'=cmd"]
    neutralised_name = [c for col in ws.iter_cols() for c in col if c.value == "'@SUM(A1:A9)"]
    assert neutralised_nr, "room nr not neutralised"
    assert neutralised_name, "room name not neutralised"
    # The numeric area cell stays a number (not stringified / not quoted).
    numeric = [c.value for col in ws.iter_cols() for c in col if c.value == 12.5]
    assert numeric, "numeric area cell must remain numeric"
