"""Audit 2026-07-28 (batch 2) — data-loss / silent-failure fixes.

Same convention as test_audit_2026_07_27_main.py: each test's docstring pins the
OLD broken behavior (reproduced before the fix) and the body asserts the fixed
one. No network calls — Vision/extractor entry points are stubbed.

DB isolation follows the house idiom: monkeypatch db._DB_PATH to a tmp file,
then init_schema() + _migrate_schema().
"""
import asyncio
import io

import pytest
from fastapi.testclient import TestClient

# main BEFORE auth — main.py's load_dotenv() must populate SAMATU_AUTH_DISABLED
# before auth.py freezes its module-level AUTH_DISABLED (see test_export_auth).
import main  # noqa: E402

import auth  # noqa: E402
import db  # noqa: E402


@pytest.fixture(autouse=True)
def _hermetic_db(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "_DB_PATH", tmp_path / "audit20260728b2.db")
    db.init_schema()
    db._migrate_schema()
    auth._ensure_dev_user()
    yield


class _FakeUpload:
    def __init__(self, data: bytes, filename: str = "plan.pdf"):
        self._data = data
        self.filename = filename

    async def read(self):
        return self._data


class _FakeRequest:
    class _Client:
        host = "127.0.0.1"

    client = _Client()


def _drain_stream(**kw):
    """Run the stream endpoint and consume the whole SSE body."""
    async def _run():
        resp = await main.extract_interior_stream_endpoint(**kw)
        return b"".join([chunk async for chunk in resp.body_iterator])
    return asyncio.run(_run())


# ══════════════════════════════════════════════════════════════════════════
# A2 — a broken lt_sanity must not truncate the SSE stream
# ══════════════════════════════════════════════════════════════════════════
# OLD: the done-event augmentation did `ls["summary"]["warnings_count"] = …`
# while the line above it used .setdefault. An lt_sanity dict with no "summary"
# key raised KeyError OUTSIDE the try that guards next(gen) → the async
# generator died → the client got a stream with NO `done` event and the UI
# spinner hung forever on an extraction that had actually SUCCEEDED.

def _patch_stream(monkeypatch, result, *, warning="PDF didelis", user_id="s1"):
    monkeypatch.delenv("BILLING_ENABLED", raising=False)
    monkeypatch.setattr(main, "_billable_user", lambda request: {"id": user_id})
    monkeypatch.setattr(main, "_check_user_quota", lambda uid: None)
    monkeypatch.setattr(main, "_check_trial_gate", lambda uid: None)
    monkeypatch.setattr(main, "_check_rate_limit", lambda ip, ep: None)
    monkeypatch.setattr(main, "_client_key", lambda request: "127.0.0.1")
    monkeypatch.setattr(main, "_validate_pdf_upload",
                        lambda b, n: {"page_count": 90, "warning": warning})
    import extractor

    def _fake_stream(pdf_bytes, overrides):
        yield {"event": "started"}
        yield {"event": "done", "result": result}
    monkeypatch.setattr(extractor, "extract_interior_stream", _fake_stream)


def test_stream_emits_done_when_lt_sanity_has_no_summary(monkeypatch):
    body = None
    _patch_stream(monkeypatch, {
        "status": "ok",
        "patalpos": [],
        # The shape that used to kill the stream: warnings present, summary absent.
        "lt_sanity": {"warnings": []},
    })
    body = _drain_stream(request=_FakeRequest(),
                         files=[_FakeUpload(b"%PDF-1.4 fake")],
                         rules_override=None).decode("utf-8")
    assert '"event": "done"' in body, "the terminal event must always ship"
    assert '"status": "ok"' in body
    # …and the counter was still maintained, via setdefault.
    assert '"warnings_count": 1' in body


def test_stream_still_emits_done_when_augmentation_raises(monkeypatch):
    """Belt-and-braces: ANY augmentation failure keeps the terminal event.

    Trigger used here is a real malformed shape, not a synthetic mock: a result
    whose `warnings` is a STRING makes `existing.append(...)` raise
    AttributeError — which, before the wrap, killed the generator exactly like
    the missing-summary KeyError did.
    """
    _patch_stream(monkeypatch, {
        "status": "ok", "patalpos": [], "warnings": "vienas įspėjimas",
    })
    body = _drain_stream(request=_FakeRequest(),
                         files=[_FakeUpload(b"%PDF-1.4 fake")],
                         rules_override=None).decode("utf-8")
    assert '"event": "done"' in body
    assert '"status": "ok"' in body


# ══════════════════════════════════════════════════════════════════════════
# B1 — a 500 must reach the browser (CORS headers on the error response)
# ══════════════════════════════════════════════════════════════════════════
# OLD: the only catch-all was @app.exception_handler(Exception), which Starlette
# lifts into ServerErrorMiddleware — OUTSIDE CORSMiddleware. The 500 body went
# out with no access-control-allow-origin, so a cross-origin fetch() rejected
# before JS could read the status: the frontend's "Nepavyko išsaugoti serveryje
# (500)" branch was unreachable and users were told it was a network problem.
# (Reproduced live on :8016 with curl before the fix; header absent.)

def test_500_response_carries_cors_headers():
    from fastapi.testclient import TestClient

    @main.app.get("/__test_boom_cors")
    async def _boom():   # pragma: no cover - the raise is the point
        raise RuntimeError("forced failure")

    try:
        c = TestClient(main.app, raise_server_exceptions=False)
        r = c.get("/__test_boom_cors", headers={"Origin": "http://localhost:5173"})
        assert r.status_code == 500
        assert r.headers.get("access-control-allow-origin") == "http://localhost:5173"
        body = r.json()
        assert body["status"] == "error"
        assert body["ref"].startswith("err-")
        # The generic LT message — never the exception text.
        assert "forced failure" not in r.text
        assert "Vidinė serverio klaida" in body["_error"]

        # A 200 on the same client is untouched.
        ok = c.get("/health", headers={"Origin": "http://localhost:5173"})
        assert ok.status_code == 200
        assert ok.headers.get("access-control-allow-origin") == "http://localhost:5173"
    finally:
        main.app.router.routes = [
            rt for rt in main.app.router.routes
            if getattr(rt, "path", None) != "/__test_boom_cors"
        ]


def test_extract_interior_summary_setdefault(monkeypatch):
    """B5-1 — the same bare-subscript bug in the NON-stream twin, where it
    500'd AFTER the paid Vision call had already produced a result."""
    monkeypatch.delenv("BILLING_ENABLED", raising=False)
    monkeypatch.setattr(main, "_billable_user", lambda request: {"id": "s2"})
    monkeypatch.setattr(main, "_check_user_quota", lambda uid: None)
    monkeypatch.setattr(main, "_check_trial_gate", lambda uid: None)
    monkeypatch.setattr(main, "_check_rate_limit", lambda ip, ep: None)
    monkeypatch.setattr(main, "_client_key", lambda request: "127.0.0.1")
    monkeypatch.setattr(main, "_validate_pdf_upload",
                        lambda b, n: {"page_count": 90, "warning": "PDF didelis"})
    import extractor
    monkeypatch.setattr(
        extractor, "extract_interior",
        lambda pdf_bytes, overrides: {
            "status": "ok", "patalpos": [], "lt_sanity": {"warnings": []},
        })

    result = asyncio.run(main.extract_interior_endpoint(
        request=_FakeRequest(), files=[_FakeUpload(b"%PDF-1.4 fake")],
        rules_override=None))
    assert result["lt_sanity"]["summary"]["warnings_count"] == 1
    assert result["_pdf_page_count"] == 90


# ══════════════════════════════════════════════════════════════════════════
# B4 — LT-comma quantities must not poison the Excel totals
# ══════════════════════════════════════════════════════════════════════════
# OLD: kiekis/kaina were written to the cell RAW. "12,5" landed as TEXT, so
# =ROUND(D*E,2) evaluated to #VALUE! and SUBTOTAL(9,…) silently SKIPPED the row
# — the client opened a sąmata whose total was short by that line's value, with
# no warning anywhere. (The formula fails loudly; the aggregate fails quietly.)

client = TestClient(main.app)


def _excel_rows_payload(rows):
    return {"projektas": "Eksporto testas", "eilutes": rows}


def _load_samata(content: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content))
    return wb, wb[wb.sheetnames[0]]


def _row_by_label(ws, label: str) -> int:
    """Locate a data row by its aprasymas — the sheet inserts section-header
    rows, so a hard-coded row index is a trap."""
    for r_ in range(1, ws.max_row + 1):
        if ws.cell(row=r_, column=2).value == label:
            return r_
    raise AssertionError(f"row {label!r} not found in the sheet")


def test_excel_lt_comma_quantity_lands_as_a_number():
    r = client.post("/generate-excel", json=_excel_rows_payload([
        {"nr": "1", "aprasymas": "Sienų dažymas", "mato_vnt": "m²",
         "kiekis": "12,5", "kaina": "7,50", "skyrius": "APDAILA"},
    ]))
    assert r.status_code == 200, r.text[:300]
    _wb, ws = _load_samata(r.content)
    row = _row_by_label(ws, "Sienų dažymas")
    kiekis, kaina = ws.cell(row=row, column=4).value, ws.cell(row=row, column=5).value
    assert isinstance(kiekis, float) and kiekis == 12.5, repr(kiekis)
    assert isinstance(kaina, float) and kaina == 7.5, repr(kaina)
    # …so the formula the row carries can actually evaluate.
    assert ws.cell(row=row, column=6).value == f"=ROUND(D{row}*E{row},2)"


def test_excel_nbsp_thousands_and_negative_price_are_numeric():
    r = client.post("/generate-excel", json=_excel_rows_payload([
        {"nr": "1", "aprasymas": "Didelis kiekis", "kiekis": "1 234,56",
         "kaina": 2},
        {"nr": "2", "aprasymas": "Nuolaida", "kiekis": 1, "kaina": "-50,00"},
    ]))
    assert r.status_code == 200, r.text[:300]
    _wb, ws = _load_samata(r.content)
    assert ws.cell(row=_row_by_label(ws, "Didelis kiekis"), column=4).value == 1234.56
    assert ws.cell(row=_row_by_label(ws, "Nuolaida"), column=5).value == -50.0


def test_excel_unparseable_quantity_stays_text_and_is_named_in_pastabos():
    r = client.post("/generate-excel", json=_excel_rows_payload([
        {"nr": "1", "aprasymas": "Neaiškus kiekis", "kiekis": "pagal poreikį",
         "kaina": 10},
    ]))
    assert r.status_code == 200, r.text[:300]
    wb, ws = _load_samata(r.content)
    row = _row_by_label(ws, "Neaiškus kiekis")
    assert ws.cell(row=row, column=4).value == "pagal poreikį"   # not silently blanked
    assert "Pastabos" in wb.sheetnames, wb.sheetnames
    notes = "\n".join(str(wb["Pastabos"].cell(row=r_, column=2).value or "")
                      for r_ in range(3, wb["Pastabos"].max_row + 1))
    assert "Neaiškus kiekis" in notes
    assert "kiekis" in notes


def test_excel_clean_numeric_payload_is_unchanged():
    r = client.post("/generate-excel", json=_excel_rows_payload([
        {"nr": "1", "aprasymas": "Dažymas", "kiekis": 10.0, "kaina": 5.0},
    ]))
    wb, ws = _load_samata(r.content)
    row = _row_by_label(ws, "Dažymas")
    assert ws.cell(row=row, column=4).value == 10.0
    assert ws.cell(row=row, column=5).value == 5.0
    assert "Pastabos" not in wb.sheetnames   # no spurious warning sheet


def test_generate_pdf_survives_an_lt_comma_quantity():
    """B5-6 — bare float('12,5') raised ValueError → 500 on the PDF export,
    for a payload the Excel export rendered fine."""
    r = client.post("/generate-pdf", json={
        "projektas": "PDF testas",
        "eilutes": [{"nr": "1", "aprasymas": "Dažymas", "mato_vnt": "m²",
                     "kiekis": "12,5", "kaina": "7,50"}],
    })
    assert r.status_code == 200, r.text[:300]
    assert r.content[:4] == b"%PDF"


def test_generate_pdf_survives_a_string_suma():
    """round('93,75') is a TypeError — it took the whole document down."""
    r = client.post("/generate-pdf", json={
        "projektas": "PDF testas",
        "eilutes": [{"nr": "1", "aprasymas": "Dažymas", "kiekis": 1,
                     "kaina": 1, "suma": "93,75"}],
    })
    assert r.status_code == 200, r.text[:300]
    assert r.content[:4] == b"%PDF"
