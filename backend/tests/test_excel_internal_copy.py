"""exports-money-5 (2026-09-30 audit) — the client workbook carried an internal
sheet.

/generate-excel added a „Pastabos" sheet of warnings meant for the contractor
(English level names, room debug formats, „Nepamirškite patikrinti", app
instructions) to the very file he sends his client. The client workbook must
hold only what the client should see; the notes sheet is kept behind an
explicit option in the export dialog (`vidine_kopija`), off by default.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from main import app  # noqa: E402
import auth  # noqa: E402


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setattr(auth, "AUTH_DISABLED", True)
    return TestClient(app)


_WARNINGS = [
    {"tekstas": "Dažomas sienų plotas 2,15× grindų ploto", "kategorija": "WARNING"},
    {"tekstas": "Svetainė (nr=05): perimetras per mažas plotui", "kategorija": "INFO · 05"},
]


def _payload(**extra):
    return {
        "projektas": "Bandomasis butas",
        "eilutes": [
            {"nr": "1", "aprasymas": "Sienų dažymas", "mato_vnt": "m²",
             "kiekis": 12.5, "kaina": 7.5, "skyrius": "APDAILA"},
        ],
        "perspejimani": list(_WARNINGS),
        **extra,
    }


def _strings(wb):
    return [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
            for c in row if isinstance(c, str)]


@pytest.mark.parametrize("url", ["/generate-excel", "/generate-excel?layout=template"])
def test_default_client_workbook_has_no_internal_sheet(client, url):
    r = client.post(url, json=_payload())
    assert r.status_code == 200, r.text[:300]
    wb = load_workbook(io.BytesIO(r.content))
    assert "Pastabos" not in wb.sheetnames, wb.sheetnames
    text = "\n".join(_strings(wb))
    assert "Nepamirškite patikrinti" not in text
    assert "WARNING" not in text and "nr=05" not in text


def test_unticked_option_is_the_default(client):
    r = client.post("/generate-excel", json=_payload(vidine_kopija=False))
    wb = load_workbook(io.BytesIO(r.content))
    assert "Pastabos" not in wb.sheetnames


def test_unparseable_number_is_not_reported_to_the_client(client):
    """The numeric-export note is for him too — it stays out of the client file
    (the cell itself still shows his text, as before)."""
    body = _payload(perspejimani=[])
    body["eilutes"][0]["kiekis"] = "pagal poreikį"
    wb = load_workbook(io.BytesIO(client.post("/generate-excel", json=body).content))
    assert "Pastabos" not in wb.sheetnames


def test_internal_copy_option_adds_the_notes_sheet(client):
    r = client.post("/generate-excel", json=_payload(vidine_kopija=True))
    assert r.status_code == 200, r.text[:300]
    wb = load_workbook(io.BytesIO(r.content))
    assert "Pastabos" in wb.sheetnames, wb.sheetnames
    notes = [wb["Pastabos"].cell(row=r_, column=2).value
             for r_ in range(3, wb["Pastabos"].max_row + 1)]
    assert notes == [w["tekstas"] for w in _WARNINGS]


def test_internal_copy_cells_still_go_through_the_safe_cell(client):
    body = _payload(
        vidine_kopija=True,
        projektas="=cmd|'/C calc'!A1",
        perspejimani=[
            {"tekstas": "=HYPERLINK(\"http://x.example\",\"x\")", "kategorija": "+1+1"},
            {"tekstas": "@SUM(A1)", "kategorija": "-2+3"},
        ],
    )
    r = client.post("/generate-excel", json=body)
    assert r.status_code == 200, r.text[:300]
    wb = load_workbook(io.BytesIO(r.content))
    assert "Pastabos" in wb.sheetnames
    # The notes sheet holds no formula of ours: every string in it is user/AI
    # text and must arrive neutralised (leading quote), none as a live formula.
    texts = [c for row in wb["Pastabos"].iter_rows(values_only=True)
             for c in row if isinstance(c, str)]
    assert any("HYPERLINK" in t for t in texts) and any("cmd|" in t for t in texts)
    for cell in texts:
        assert cell[:1] not in ("=", "+", "-", "@"), cell
