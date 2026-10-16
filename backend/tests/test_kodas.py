"""Cost codes („kodas", togal-v2 W9): one rule, server and browser alike.

The cases live in frontend/src/lib/kodas.cases.json and are run by
frontend/src/lib/kodas.test.js too. Plus where the code travels: the optional
Excel column, the F2 CSV „Kodas" column, the samata-render rows.
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import auth  # noqa: E402
from api.sudetis import effective_kodas  # noqa: E402
from main import app  # noqa: E402

CASES = json.loads((ROOT.parent / "frontend" / "src" / "lib" / "kodas.cases.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("case", CASES["cases"], ids=[c["name"] for c in CASES["cases"]])
def test_effective_kodas_shared_cases(case):
    assert effective_kodas(case["row"], CASES["own"]) == case["want"]


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setattr(auth, "AUTH_DISABLED", True)
    return TestClient(app)


def _payload(**extra):
    return {"projektas": "Bandomasis butas", "eilutes": [
        {"nr": "1", "aprasymas": "Sienų glaistymas/dažymas", "mato_vnt": "m²", "kiekis": 100.0,
         "kaina": 5.0, "skyrius": "Sienų glaistymas dažymas", "kodas": "01.01"},
        {"nr": "2", "aprasymas": "Klozeto montavimas", "mato_vnt": "vnt", "kiekis": 2.0,
         "kaina": 40.0, "skyrius": "Santechnikos darbai", "kodas": "06"},
    ], **extra}


def _main(r):
    assert r.status_code == 200, r.text[:300]
    return load_workbook(io.BytesIO(r.content))["Sąmata"]


@pytest.mark.parametrize("grouping", [None, "patalpos", "be_grupavimo"])
def test_excel_kodas_column_is_optional_and_off_by_default(client, grouping):
    extra = {"grupavimas": grouping} if grouping else {}
    off = _main(client.post("/generate-excel", json=_payload(**extra)))
    assert "Kodas" not in [c.value for c in off[3]]
    on = _main(client.post("/generate-excel", json=_payload(su_kodu=True, **extra)))
    head = [c.value for c in on[3]]
    assert "Kodas" in head
    col = head.index("Kodas") + 1
    codes = {on.cell(row=r, column=2).value: on.cell(row=r, column=col).value
             for r in range(4, on.max_row + 1)}
    assert codes["Sienų glaistymas/dažymas"] == "01.01"
    assert codes["Klozeto montavimas"] == "06"
    # the money does not move with a column
    assert [c.value for c in on["F"]] == [c.value for c in off["F"]]


def test_f2_lines_print_the_cost_code_not_the_row_id(monkeypatch):
    from api import exports_v2
    project = {"user_id": None,
               "pozicijos": [{"id": "abc123", "skyrius": "Sienų glaistymas dažymas",
                              "aprasymas": "Sienų glaistymas/dažymas", "vnt": "m2", "kiekis": 10},
                             {"id": "def456", "skyrius": "Santechnikos darbai",
                              "aprasymas": "Klozeto montavimas", "vnt": "vnt", "kiekis": 2}],
               "darbu_sarasas": [{"id": "def456", "kodas": "S-9", "skyrius_enum": "plumbing"}]}
    monkeypatch.setattr(exports_v2.db, "get_project", lambda pid: project)
    monkeypatch.setattr(exports_v2, "PROJECTS_DIR", Path("/nonexistent-w9"))
    lines = exports_v2._load_project_lines("p1")
    assert [ln["code"] for ln in lines] == ["01.01", "S-9"]


def test_samata_render_rows_carry_the_shared_code(monkeypatch):
    from api import samata_templates
    monkeypatch.setattr(samata_templates._db, "list_assemblies", lambda user_id: [])
    rows_in = [{"id": "a", "aprasymas": "Sienų glaistymas/dažymas", "skyrius_enum": "paint"},
               {"id": "b", "aprasymas": "Sienų glaistymas/dažymas virš 3 m", "skyrius_enum": "paint"},
               {"id": "c", "aprasymas": "Klozeto montavimas", "skyrius_enum": "plumbing"}]
    rows = [{"_saltinio_id": ["a", "b"]}, {"_saltinio_id": ["a", "c"]}, {"_saltinio_id": []}]
    samata_templates._attach_kodai(rows, rows_in, {"id": "u"})
    assert rows[0]["kodas"] == "01.01"
    assert "kodas" not in rows[1] and "kodas" not in rows[2]
