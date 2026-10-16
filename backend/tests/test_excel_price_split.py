"""The darbo + medžiagų split prints only in his internal copy (togal-v2 W9).

A rinkinys (or a kainynas material column) splits a price into kaina_darbo +
kaina_medziagu. The client's workbook never shows the split; the internal copy
adds two numeric columns after „Suma €", in every layout, only when some line
carries one. `suma` and the totals do not change either way.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from main import app  # noqa: E402
import auth  # noqa: E402


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setattr(auth, "AUTH_DISABLED", True)
    return TestClient(app)


def _payload(split=True, **extra):
    row = {"nr": "1", "aprasymas": "Sienų glaistymas/dažymas", "mato_vnt": "m²",
           "kiekis": 100.0, "kaina": 5.3, "skyrius": "Sienų glaistymas dažymas"}
    if split:
        row.update(kaina_darbo=4.2, kaina_medziagu=1.1)
    return {"projektas": "Bandomasis butas", "eilutes": [
        row,
        {"nr": "2", "aprasymas": "Grindjuosčių sumontavimas", "mato_vnt": "m",
         "kiekis": 10.0, "kaina": 3.0, "skyrius": "Sienų glaistymas dažymas"},
    ], **extra}


def _main(r):
    assert r.status_code == 200, r.text[:300]
    return load_workbook(io.BytesIO(r.content))["Sąmata"]


def _header(ws):
    return [c.value for c in ws[3]]


def _cells(ws):
    return [c for row in ws.iter_rows(values_only=True) for c in row]


@pytest.mark.parametrize("grouping", [None, "patalpos", "be_grupavimo"])
def test_internal_copy_shows_the_split_as_numbers(client, grouping):
    extra = {"vidine_kopija": True}
    if grouping:
        extra["grupavimas"] = grouping
    ws = _main(client.post("/generate-excel", json=_payload(**extra)))
    head = _header(ws)
    assert "Darbo kaina €" in head and "Medžiagų kaina €" in head
    d, m = head.index("Darbo kaina €") + 1, head.index("Medžiagų kaina €") + 1
    row = next(r for r in range(4, ws.max_row + 1)
               if ws.cell(row=r, column=2).value == "Sienų glaistymas/dažymas")
    assert ws.cell(row=row, column=d).value == 4.2
    assert ws.cell(row=row, column=m).value == 1.1


@pytest.mark.parametrize("grouping", [None, "patalpos"])
def test_client_copy_never_shows_the_split(client, grouping):
    extra = {"grupavimas": grouping} if grouping else {}
    ws = _main(client.post("/generate-excel", json=_payload(**extra)))
    assert "Darbo kaina €" not in _header(ws) and "Medžiagų kaina €" not in _header(ws)


def test_internal_copy_without_any_split_is_unchanged(client):
    a = _main(client.post("/generate-excel", json=_payload(split=False, vidine_kopija=True)))
    b = _main(client.post("/generate-excel", json=_payload(split=False)))
    assert _header(a) == _header(b)
    assert _cells(a) == _cells(b)


def test_totals_do_not_move_with_the_split(client):
    a = _main(client.post("/generate-excel", json=_payload(vidine_kopija=True)))
    b = _main(client.post("/generate-excel", json=_payload(split=False)))
    assert [c for c in _cells(a) if isinstance(c, str) and c.startswith("=")] == \
           [c for c in _cells(b) if isinstance(c, str) and c.startswith("=")]
