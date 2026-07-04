"""Tests for /api/projects/{id}/export-excel.

P2 (2026-07-04 pentest, pass 3): the endpoint now requires auth + URL-project
ownership BEFORE the ``?fixture=`` branch (see api/export_excel._load_project).
These functional tests therefore log in as the project owner; the P2 IDOR PoC
(anon + fixture → XLSX leaking another tenant's calibration) is replayed in
test_p2_export_excel_fixture_auth.py.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "backend"))

_PW = "passw0rd123"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    """A TestClient logged in as the owner of the projects used below."""
    from main import app
    import auth
    import db

    monkeypatch.setattr(db, "_DB_PATH", tmp_path / "export.db")
    monkeypatch.setattr(auth, "AUTH_DISABLED", False)
    db.init_schema()
    db._migrate_schema()
    owner = db.create_user("owner@x.lt", auth.hash_password(_PW), name="Owner")
    for pid in ("test", "no-calibration-yet"):
        db.upsert_project(pid, {"name": pid, "tipas": "interjeras"}, user_id=owner["id"])

    c = TestClient(app)
    r = c.post("/api/auth/login", json={"email": "owner@x.lt", "password": _PW})
    assert r.status_code == 200, r.text
    return c


def test_export_returns_xlsx_with_ai_sheet(client):
    r = client.post("/api/projects/test/export-excel?fixture=berzas")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    assert "AI Sąmata" in wb.sheetnames
    ws = wb["AI Sąmata"]
    # Header row should exist
    assert ws["A3"].value == "Kategorija"
    assert ws["B3"].value == "AI"
    assert ws["C3"].value == "Patikslinta"


def test_export_ai_sheet_carries_real_numbers(client):
    r = client.post("/api/projects/test/export-excel?fixture=berzas")
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    ws = wb["AI Sąmata"]
    # Find the dazymo row by label and check the AI column has a number > 0
    found = False
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] == "Dažymo plotas (visas)":
            assert isinstance(row[1], (int, float))
            assert row[1] > 0
            found = True
            break
    assert found, "Dažymo plotas row missing"


def test_export_includes_per_room_block(client):
    r = client.post("/api/projects/test/export-excel?fixture=berzas")
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    ws = wb["AI Sąmata"]
    # The per-room block has header "Patalpos"; scan for it
    found_block_header = False
    for row in ws.iter_rows(values_only=True):
        if row[0] == "Patalpos":
            found_block_header = True
            break
    assert found_block_header


def test_export_without_calibration_has_no_second_sheet(client):
    r = client.post("/api/projects/no-calibration-yet/export-excel?fixture=berzas")
    wb = load_workbook(io.BytesIO(r.content), data_only=True)
    assert wb.sheetnames == ["AI Sąmata"]


def test_unknown_fixture_returns_404(client):
    r = client.post("/api/projects/test/export-excel?fixture=does-not-exist")
    assert r.status_code == 404
