"""A made-up price list, uploaded end to end (founder request, 2026-09-30).

„Make a fake price-list document, upload it, and make sure the app pulls the
prices out correctly." The list lives in ``fixtures_price_list/`` as a spec
(``kainynas_eilutes.json``) with the expected result for every row; this file
builds the .xlsx from it, sends it through the real upload endpoint and checks
row by row. The same spec drives the browser-side CSV test
(``frontend/src/lib/priceListCsv.test.js``).

Every name in the list is invented. Run this file directly to rewrite the two
files a person can upload by hand:

    ../venv/bin/python tests/test_price_list_made_up.py
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import api.price_lists as pl_mod   # noqa: E402
import auth as _auth               # noqa: E402

FIXTURES = HERE / "fixtures_price_list"
SPEC = json.loads((FIXTURES / "kainynas_eilutes.json").read_text(encoding="utf-8"))
XLSX_PATH = FIXTURES / "bandomasis_kainynas.xlsx"
CSV_PATH = FIXTURES / "bandomasis_kainynas.csv"

# Title on row 1, a blank row 2, the header on row 3 — the way a contractor's
# own sheet starts. Data rows therefore begin on Excel row 4.
FIRST_DATA_ROW = 4


def build_xlsx(spec: dict = SPEC) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kainynas"
    ws.append([spec["title"]])
    ws.append([])
    ws.append(spec["header"])
    for row in spec["rows"]:
        # A JSON number becomes a real Excel number cell, a string a text cell.
        ws.append([row["darbas"], row["vnt"] or None, row["kaina"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(spec: dict = SPEC) -> str:
    """What „Excel“ writes for „CSV UTF-8“ in a Lithuanian locale: a BOM,
    `;` between columns, the decimal comma."""
    lines = [";".join(spec["header"])]
    for row in spec["rows"]:
        lines.append(";".join([row["darbas"], row["vnt"], row["csv"]]))
    return "﻿" + "\r\n".join(lines) + "\r\n"


def _imported_rows(spec: dict = SPEC) -> list[dict]:
    return [r for r in spec["rows"] if "kaina" in r["expect"]]


def _unreadable_rows(spec: dict = SPEC) -> list[tuple[int, dict]]:
    return [(FIRST_DATA_ROW + i, r) for i, r in enumerate(spec["rows"])
            if r["expect"].get("skip") == "unreadable"]


def _check_against_spec(body: dict) -> None:
    want = _imported_rows()
    got = body["items"]
    assert [i["aprasymas"] for i in got] == [r["darbas"] for r in want]
    for item, row in zip(got, want):
        exp = row["expect"]
        assert item["kaina"] == pytest.approx(exp["kaina"]), \
            f"{row['darbas']}: cell {row['kaina']!r} gave {item['kaina']}, expected {exp['kaina']}"
        assert item["mato_vnt"] == exp["mato_vnt"], row["darbas"]
        assert item["skyrius"] == exp["skyrius"], row["darbas"]
    skipped = [r for r in SPEC["rows"] if "skip" in r["expect"]]
    assert body["skipped_rows"] == len(skipped)
    assert body["problem_rows"] == [
        {"row": n, "aprasymas": r["darbas"], "kaina": r["kaina"]}
        for n, r in _unreadable_rows()
    ]


@pytest.fixture()
def client(monkeypatch):
    app = FastAPI()
    app.include_router(pl_mod.router)
    monkeypatch.setattr(_auth, "require_user",
                        lambda request: {"id": "made-up-price-list-user", "is_admin": False})
    return TestClient(app)


def test_the_made_up_price_list_uploads_through_the_real_endpoint(client):
    r = client.post("/api/price-lists/parse-xlsx",
                    files={"file": ("bandomasis_kainynas.xlsx", build_xlsx())})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["header_row"] == FIRST_DATA_ROW - 1
    _check_against_spec(body)


def test_the_files_for_a_hand_upload_match_the_spec(client):
    """The .xlsx and .csv beside the spec are what a person uploads by hand;
    they must say exactly what the spec says."""
    assert CSV_PATH.read_bytes().decode("utf-8") == build_csv(), \
        "run this file directly to rewrite the fixture files"
    r = client.post("/api/price-lists/parse-xlsx",
                    files={"file": (XLSX_PATH.name, XLSX_PATH.read_bytes())})
    assert r.status_code == 200, r.text
    _check_against_spec(r.json())


# ── the price cell on its own ────────────────────────────────────────────────

@pytest.mark.parametrize("cell, price, unit", [
    ("15", 15.0, None),
    ("15,50", 15.5, None),
    ("15.50", 15.5, None),
    ("15 €", 15.0, None),
    ("15€", 15.0, None),
    ("15 €/m2", 15.0, "m2"),          # was 152
    ("5 Eur/m2", 5.0, "m2"),          # was 52
    ("12,50 €/m2", 12.5, "m2"),       # was 12.502
    ("12,5 Eur/m²", 12.5, "m²"),      # was refused
    ("8 €/m3", 8.0, "m3"),            # was 83
    ("3 € už m3", 3.0, "m3"),         # was 33
    ("4,10 EUR / m", 4.1, "m"),
    ("45 € vnt.", 45.0, "vnt."),
    ("120 eur/kompl.", 120.0, "kompl."),
    ("7 m²", 7.0, "m²"),
    ("1 234,56 €", 1234.56, None),
    ("1 234,56 €", 1234.56, None),   # non-breaking space
    ("1 234,56", 1234.56, None),     # narrow no-break space (what „Excel“ emits)
    ("1 234,56", 1234.56, None),     # thin space
    ("1.234,56", 1234.56, None),
    ("1,234.56", 1234.56, None),
    (7.5, 7.5, None),
    (12, 12.0, None),
])
def test_a_price_cell_reads_the_leading_number_and_its_unit(cell, price, unit):
    got_price, got_unit, status = pl_mod._read_price_cell(cell)
    assert status == "ok"
    assert got_price == pytest.approx(price)
    assert got_unit == unit


@pytest.mark.parametrize("cell", [
    "sutartinė",
    "15–20 €",            # a range is not a price
    "15-20",
    "15 € be pagrindo",   # words after the price
    "15 2",               # a space that is not a thousands separator
    "apie 15 €",          # the number is not first
    "vnt.",
    "-",
])
def test_anything_else_in_a_price_cell_is_unreadable_not_a_number(cell):
    price, unit, status = pl_mod._read_price_cell(cell)
    assert (price, unit, status) == (None, None, "unreadable")


@pytest.mark.parametrize("cell", [None, "", "   "])
def test_an_empty_price_cell_is_empty(cell):
    assert pl_mod._read_price_cell(cell) == (None, None, "empty")


def test_a_date_that_excel_made_out_of_a_price_is_not_a_price():
    """„15.05“ typed into a cell can come back as a date; its digits must never
    be glued into a number."""
    import datetime as _dt
    assert pl_mod._read_price_cell(_dt.datetime(2026, 5, 15))[2] == "unreadable"


def test_an_unreadable_price_is_skipped_and_named_not_imported():
    wb = Workbook()
    ws = wb.active
    ws.append(["Darbas", "Mato vnt.", "Kaina"])
    ws.append(["Sienų dažymas", "m²", "15 €/m2"])
    ws.append(["Durų montavimas", "vnt.", "sutartinė"])
    buf = io.BytesIO()
    wb.save(buf)
    out = pl_mod.parse_xlsx_items(buf.getvalue())
    assert [i["kaina"] for i in out["items"]] == [15.0]
    assert out["skipped_rows"] == 1
    assert out["problem_rows"] == [
        {"row": 3, "aprasymas": "Durų montavimas", "kaina": "sutartinė"}]


def test_nothing_readable_is_still_a_400():
    wb = Workbook()
    ws = wb.active
    ws.append(["Darbas", "Kaina"])
    ws.append(["Sienų dažymas", "sutartinė"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(HTTPException) as exc:
        pl_mod.parse_xlsx_items(buf.getvalue())
    assert exc.value.status_code == 400


if __name__ == "__main__":
    XLSX_PATH.write_bytes(build_xlsx())
    CSV_PATH.write_text(build_csv(), encoding="utf-8", newline="")
    print(f"wrote {XLSX_PATH}\nwrote {CSV_PATH}")
