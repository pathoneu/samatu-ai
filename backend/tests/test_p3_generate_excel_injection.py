"""P3 (LOW) — CSV/Excel formula injection on the /generate-excel offer sheet.

2026-07-04 pentest, pass 3, finding P3 (F6 sibling). F6 (78041cd) neutralised the
sibling api/export_excel.py + f2_export.py but MISSED main.py's /generate-excel
handler, which writes user/AI text (imone, projektas, skyrius, nr, aprašymas,
mato_vnt) straight into openpyxl cells. A leading = + - @ would evaluate as a
live formula when the client-facing .xlsx is opened.

The fix wraps those text/label cells (never the numeric kiekis/kaina cells) with
extractor.spreadsheet_safety.sanitize_spreadsheet_cell. These tests REPLAY the
PoC and assert every risky value is neutralised with a leading single quote while
numbers stay numeric.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from main import app  # noqa: E402
import auth  # noqa: E402


@pytest.fixture()
def client(monkeypatch):
    # /generate-excel only needs require_user; run under the dev bypass so the
    # test exercises the WRITER path (the injection fix), not the auth gate.
    monkeypatch.setattr(auth, "AUTH_DISABLED", True)
    return TestClient(app)


def _all_string_cells(content: bytes):
    wb = load_workbook(io.BytesIO(content), data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, str):
                    out.append(cell)
    return out, wb


def test_p3_generate_excel_neutralises_all_injection_vectors(client):
    payload = {
        "projektas": "=cmd|'/C calc'!A1",           # user title → A1 banner
        "imone": '=HYPERLINK("http://evil.example","x")',  # user co. → A1 banner
        "eilutes": [
            {
                "skyrius": "=SUM(A1:A9)",            # AI section header
                "nr": "-2+3",                        # AI text cell
                "aprasymas": "=cmd|calc",            # AI text cell
                "mato_vnt": "@SUM(A1)",              # AI text cell
                "kiekis": 12.5,                      # numeric — must stay numeric
                "kaina": 4.0,                        # numeric — must stay numeric
            },
        ],
    }
    r = client.post("/generate-excel", json=payload)
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]

    strings, wb = _all_string_cells(r.content)

    # Every attacker-controlled text cell is present ONLY in its neutralised form.
    assert any(s.startswith("'=HYPERLINK") for s in strings), "imone banner not neutralised"
    assert any(s == "'=SUM(A1:A9)" for s in strings), "skyrius not neutralised"
    assert any(s == "'-2+3" for s in strings), "nr not neutralised"
    assert any(s == "'=cmd|calc" for s in strings), "aprašymas not neutralised"
    assert any(s == "'@SUM(A1)" for s in strings), "mato_vnt not neutralised"

    # No live-formula lead survives on ANY string cell.
    for s in strings:
        assert s[0] not in ("=", "+", "-", "@", "\t", "\r", "\n"), \
            f"unneutralised formula-trigger cell survived: {s!r}"

    # Numeric quantity / price cells stay numeric (not quoted / not stringified).
    ws = wb.active
    numeric_vals = [c for row in ws.iter_rows(values_only=True) for c in row
                    if isinstance(c, (int, float))]
    assert any(abs(v - 12.5) < 1e-6 for v in numeric_vals), "kiekis must stay numeric"
    assert any(abs(v - 4.0) < 1e-6 for v in numeric_vals), "kaina must stay numeric"
