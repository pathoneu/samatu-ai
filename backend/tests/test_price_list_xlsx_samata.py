"""A contractor's own sąmata, uploaded as a kainynas (2026-09-15).

The founder's test: take an old sąmata .xlsx and upload it as a price list.
Such a workbook has a title block above the header, a two-row header
(„Darbų kaina € be PVM" over „vnt. | viso"), section rows („Skyrius | …")
and a materials price beside the work price. The parser used to look at the
first five rows only and knew the bare words „aprašymas" / „kaina" — the
upload failed with „Nerasta antraštės eilutė". Synthetic rows only.
"""
from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi import HTTPException

from api.price_lists import parse_xlsx_items


def _samata_workbook(header_at: int = 8) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "UAB Pavyzdys"
    ws["A2"] = "Objektas: butas"
    ws["A4"] = "Sąmata Nr. 7"
    r = header_at
    ws.cell(r, 1, "Eil. Nr."); ws.cell(r, 2, "Darbų aprašymai"); ws.cell(r, 3, "Mato Vnt."); ws.cell(r, 4, "Kiekis")
    ws.cell(r, 5, "Darbų kaina € be PVM"); ws.cell(r, 7, "Medžiagų kaina € be PVM")
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
    ws.cell(r + 1, 5, "vnt."); ws.cell(r + 1, 6, "viso"); ws.cell(r + 1, 7, "vnt."); ws.cell(r + 1, 8, "viso")
    ws.cell(r + 2, 1, "Skyrius"); ws.cell(r + 2, 2, "Gipso kartono montavimas")
    ws.cell(r + 3, 1, 1); ws.cell(r + 3, 2, "GKP montavimas ant lubų"); ws.cell(r + 3, 3, "m2"); ws.cell(r + 3, 4, 137.71)
    ws.cell(r + 3, 5, 15); ws.cell(r + 3, 6, 2065.65); ws.cell(r + 3, 7, 18); ws.cell(r + 3, 8, 2478.78)
    ws.cell(r + 4, 1, 2); ws.cell(r + 4, 2, "GKP siūlių glaistymas"); ws.cell(r + 4, 3, "m"); ws.cell(r + 4, 4, 31.26)
    ws.cell(r + 4, 5, "11,50"); ws.cell(r + 4, 6, 359.49); ws.cell(r + 4, 7, 0); ws.cell(r + 4, 8, 0)
    ws.cell(r + 6, 2, "Plytelių klijavimas")                      # a lone section name, no price
    ws.cell(r + 7, 1, 3); ws.cell(r + 7, 2, "Grindų plytelių klijavimas"); ws.cell(r + 7, 3, "m2"); ws.cell(r + 7, 4, 19.35)
    ws.cell(r + 7, 5, 18); ws.cell(r + 7, 6, 348.3); ws.cell(r + 7, 7, 20); ws.cell(r + 7, 8, 387)
    ws.cell(r + 9, 2, "Iš viso:"); ws.cell(r + 9, 6, 2773.44)      # a totals row: description without a unit price
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_a_contractors_samata_uploads_as_a_price_list():
    out = parse_xlsx_items(_samata_workbook())
    assert out["header_row"] == 8
    items = out["items"]
    assert [i["aprasymas"] for i in items] == ["GKP montavimas ant lubų", "GKP siūlių glaistymas", "Grindų plytelių klijavimas"]
    assert [i["kaina"] for i in items] == [15.0, 11.5, 18.0], "the WORK unit price, never the total"
    assert [i["mato_vnt"] for i in items] == ["m2", "m", "m2"]
    assert [i["skyrius"] for i in items] == ["Gipso kartono montavimas", "Gipso kartono montavimas", "Plytelių klijavimas"]
    assert items[0]["medziagu_kaina"] == 18.0 and "medziagu_kaina" not in items[1]
    # the sub-header, the section rows and the totals row are skipped, not imported
    assert out["skipped_rows"] == 4


def test_a_plain_kainynas_still_parses_exactly_as_before():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["skyrius", "aprašymas", "vnt", "kaina"])
    ws.append(["Apdaila", "Sienų dažymas", "m2", 4.5])
    ws.append(["Apdaila", "Lubų dažymas", "m2", "5,00 €"])
    ws.append(["", "Be kainos", "m2", None])
    buf = io.BytesIO(); wb.save(buf)
    out = parse_xlsx_items(buf.getvalue())
    assert out["header_row"] == 1 and out["skipped_rows"] == 1
    assert [(i["skyrius"], i["aprasymas"], i["kaina"]) for i in out["items"]] == [
        ("Apdaila", "Sienų dažymas", 4.5), ("Apdaila", "Lubų dažymas", 5.0)]


def test_a_header_too_deep_or_missing_is_still_refused_in_lithuanian():
    with pytest.raises(HTTPException) as exc:
        parse_xlsx_items(_samata_workbook(header_at=30))
    assert exc.value.status_code == 400 and "antraštės" in exc.value.detail
