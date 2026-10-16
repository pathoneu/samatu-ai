"""/generate-excel with „Grupuoti pagal" (togal-v2 W8): the grouped main sheet prints the same quantities and
the same total as today's layout — only order and subtotals change. No grouping = today's workbook."""
from __future__ import annotations

import io
import re
from decimal import ROUND_HALF_UP, Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app
from tests.test_export_grouping import ROOMS, SHEETS, _lines

client = TestClient(app)


def _payload(grupavimas=None, **extra):
    p = {"projektas": "Testas", "tipas": "interjeras", "eilutes": _lines(), "patalpos": ROOMS,
         "lapai": SHEETS, **extra}
    if grupavimas is not None:
        p["grupavimas"] = grupavimas
    return p


def _cents(x) -> Decimal:
    return Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _main_sheet(content):
    return load_workbook(io.BytesIO(content))["Sąmata"]


def _data_rows(ws):
    """(nr, aprašymas, kiekis, suma as Decimal, extra G…) for every printed line row of the main sheet."""
    out = []
    for r in range(4, ws.max_row + 1):
        d, e, f = ws.cell(r, 4).value, ws.cell(r, 5).value, ws.cell(r, 6).value
        if not isinstance(d, (int, float)):
            continue
        if isinstance(f, str) and re.fullmatch(rf"=ROUND\(D{r}\*E{r},2\)", f):
            amount = _cents(d * (e or 0))
        else:
            amount = Decimal(str(f))
        extra = [ws.cell(r, c).value for c in range(7, 12)]
        out.append((str(ws.cell(r, 1).value), ws.cell(r, 2).value, d, amount, extra))
    return out


def test_no_grouping_is_todays_workbook():
    a = client.post("/generate-excel", json=_payload())
    b = client.post("/generate-excel", json=_payload("skyriai"))
    assert a.status_code == 200 and b.status_code == 200
    ra, rb = _data_rows(_main_sheet(a.content)), _data_rows(_main_sheet(b.content))
    assert ra == rb
    ws = _main_sheet(a.content)
    assert ws.cell(3, 7).value is None                     # no extra columns in today's layout
    assert all(isinstance(ws.cell(r, 6).value, str) for r in range(4, ws.max_row + 1)
               if isinstance(ws.cell(r, 4).value, (int, float)))  # Suma stays a live formula


@pytest.mark.parametrize("mode", ["patalpos", "seimos", "lapai", "be_grupavimo"])
def test_grouped_sheet_prints_the_same_lines_and_total(mode):
    base = _data_rows(_main_sheet(client.post("/generate-excel", json=_payload()).content))
    r = client.post("/generate-excel", json=_payload(mode))
    assert r.status_code == 200, r.text
    ws = _main_sheet(r.content)
    rows = _data_rows(ws)
    assert sum(x[3] for x in rows) == sum(x[3] for x in base)          # money to the cent
    per_line: dict[str, float] = {}
    for nr, _a, kiekis, _s, _e in rows:
        per_line[nr] = per_line.get(nr, 0.0) + kiekis
    for nr, _a, kiekis, _s, _e in base:
        assert per_line[nr] == pytest.approx(kiekis, abs=1e-9), (mode, nr)
    # the grand total is the same SUBTOTAL over the rows; every row says how it was obtained
    heads = [ws.cell(3, c).value for c in range(7, 11)]
    assert "Kaip gauta" in heads
    col = 7 + heads.index("Kaip gauta")
    kinds = {ws.cell(rr, col).value for rr in range(4, ws.max_row + 1) if isinstance(ws.cell(rr, 4).value, (int, float))}
    assert "Pataisyta jūsų" in kinds and "Programa — skaityta tiesiogiai" in kinds


def test_rooms_grouping_prints_room_headings_and_room_subtotals():
    ws = _main_sheet(client.post("/generate-excel", json=_payload("patalpos")).content)
    texts = [ws.cell(r, 1).value for r in range(4, ws.max_row + 1)]
    assert "01 Tambūras" in texts and "Nepriskirta patalpoms" in texts
    assert "Patalpos suma:" in texts
    assert "pagal patalpas" in ws["A1"].value


def test_flat_export_has_filter_and_attribute_columns():
    ws = _main_sheet(client.post("/generate-excel", json=_payload("be_grupavimo")).content)
    assert [ws.cell(3, c).value for c in range(7, 11)] == ["Skyrius", "Patalpa", "Lapas", "Kaip gauta"]
    assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A3:J")
    texts = [ws.cell(r, 1).value for r in range(4, ws.max_row + 1)]
    assert "Skyriaus suma:" not in texts                   # no headings, no subtotals


def test_assumption_column_only_in_the_internal_copy():
    lines = _lines()
    lines[0]["prielaida"] = "Prielaida: 30 % sienų lyginama"
    client_copy = _main_sheet(client.post("/generate-excel", json={**_payload("patalpos"), "eilutes": lines}).content)
    assert "Prielaida" not in [client_copy.cell(3, c).value for c in range(7, 12)]
    internal = _main_sheet(client.post(
        "/generate-excel", json={**_payload("patalpos", vidine_kopija=True), "eilutes": lines}).content)
    assert "Prielaida" in [internal.cell(3, c).value for c in range(7, 12)]


def test_unknown_grouping_is_a_400():
    r = client.post("/generate-excel", json=_payload("bet-kas"))
    assert r.status_code == 400


def test_his_layout_ignores_the_grouping():
    r = client.post("/generate-excel?layout=template", json=_payload("patalpos"))
    assert r.status_code == 200
    ws = _main_sheet(r.content)
    assert ws.cell(3, 7).value is None
