"""Re-audit of the 2026-09-29 security findings (v3 W7-A M3), one test per fix.

  -4   an unparseable kiekis/kaina is text in the workbook (never a live
       formula) and plain text in the PDF (never markup)
  -5   material-specs: rate bucket + spend cap + reservation before a paid read
  -10  room_recalc ?fixture= is ignored in production (auth first)
  -11  telemetry/session needs a login, an owned project, bounded input
  -12  auth responses carry no google_sub / Stripe ids / session nonce
  -21  calibrate returns no server path; v9 symbol errors carry no exception text
"""
from __future__ import annotations

import io
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app  # noqa: E402

import auth  # noqa: E402
import db  # noqa: E402

_PW = "passw0rd123"


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "_DB_PATH", tmp_path / "reaudit.db")
    monkeypatch.setattr(auth, "AUTH_DISABLED", False)
    db.init_schema()
    db._migrate_schema()
    owner = db.create_user("owner@x.lt", auth.hash_password(_PW), name="Owner")
    other = db.create_user("other@x.lt", auth.hash_password(_PW), name="Other")
    db.upsert_project("p-own", {"name": "Mano", "tipas": "interjeras",
                                "patalpos": [{"nr": "1", "pavadinimas": "Svetainė"}],
                                "puslapiai": []}, user_id=owner["id"])
    db.upsert_project("p-other", {"name": "Kito", "tipas": "interjeras"}, user_id=other["id"])
    return {"owner": owner, "other": other}


def _login(email):
    c = TestClient(app)
    assert c.post("/api/auth/login", json={"email": email, "password": _PW}).status_code == 200
    return c


# ── -12 ──────────────────────────────────────────────────────────────────────

def test_auth_responses_carry_no_private_user_fields(env):
    db.set_stripe_customer_id(env["owner"]["id"], "cus_secret")
    c = TestClient(app)
    login = c.post("/api/auth/login", json={"email": "owner@x.lt", "password": _PW}).json()["user"]
    me = c.get("/api/auth/me").json()["user"]
    reg = TestClient(app).post("/api/auth/register",
                               json={"email": "new@x.lt", "password": _PW}).json()["user"]
    for u in (login, me, reg):
        for k in ("password_hash", "google_sub", "stripe_customer_id", "stripe_subscription_id",
                  "session_version", "deleted_at"):
            assert k not in u, (k, u)
        assert u["email"] and "id" in u and "is_admin" in u


# ── -21 ──────────────────────────────────────────────────────────────────────

def test_calibrate_returns_no_server_path(env):
    c = _login("owner@x.lt")
    r = c.post("/api/projects/p-own/calibrate", json={"overrides": {}})
    assert r.status_code == 200, r.text
    assert "stored_at" not in r.json()
    assert "/" not in str(r.json().get("factors", ""))


def test_v9_symbol_errors_do_not_echo_exceptions(env):
    import main
    out = main._v9_error(RuntimeError("/srv/app/backend/interior_v9.py: secret detail"), "run")
    assert out["status"] == "error"
    assert "secret" not in out["error"] and "/srv" not in out["error"]
    assert "klaidos nr." in out["error"]


# ── -4 ───────────────────────────────────────────────────────────────────────

def test_unparseable_quantity_is_text_in_excel(env):
    c = _login("owner@x.lt")
    payload = {"projektas": "P", "eilutes": [
        {"nr": "1", "aprasymas": "Dažymas", "mato_vnt": "m²", "kiekis": "=HYPERLINK(\"http://x\")",
         "kaina": "@SUM(A1)", "suma": 0, "skyrius": "Dažymas"}]}
    r = c.post("/generate-excel", json=payload)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    texts = [v for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row
             if isinstance(v, str)]
    assert not any(t.startswith(("=HYPERLINK", "@SUM")) for t in texts), texts
    assert any(t.startswith("'=HYPERLINK") for t in texts)


def test_unparseable_quantity_is_text_in_pdf(env):
    c = _login("owner@x.lt")
    payload = {"projektas": "P", "eilutes": [
        {"nr": "1", "aprasymas": "Dažymas", "mato_vnt": "m²", "kiekis": "<font size=80>9</font><b",
         "kaina": 3, "suma": 0, "skyrius": "Dažymas"}]}
    r = c.post("/generate-pdf", json=payload)
    assert r.status_code == 200, r.text
    assert r.content[:4] == b"%PDF"


# ── -5 ───────────────────────────────────────────────────────────────────────

def test_material_specs_is_capped_and_reserved_after_ownership(env, monkeypatch):
    import main
    calls = []
    monkeypatch.setattr(main, "_pick_spec_pages_for_room", lambda room, pages: [1])
    monkeypatch.setattr(main, "_check_user_spend_cap", lambda uid: calls.append(("cap", uid)))
    monkeypatch.setattr(main, "_reserve_billable", lambda uid, path: calls.append(("reserve", path)))
    monkeypatch.setattr(main, "_validate_pdf_upload", lambda b, n: None)
    c = _login("owner@x.lt")
    files = {"files": ("a.pdf", b"%PDF-1.4 x", "application/pdf")}
    foreign = c.post("/api/projects/p-other/material-specs", files=files, data={"room_nr": "1"})
    assert foreign.status_code == 404 and calls == []          # ownership first, nothing charged
    with patch("extractor.material_specs.extract_room_material_specs", return_value={"paint": None}):
        r = c.post("/api/projects/p-own/material-specs", files=files, data={"room_nr": "1"})
    assert r.status_code == 200, r.text
    assert ("cap", env["owner"]["id"]) in calls
    assert ("reserve", "/api/projects/p-own/material-specs") in calls


# ── -10 ──────────────────────────────────────────────────────────────────────

def test_room_recalc_fixture_ignored_in_production(env, monkeypatch):
    monkeypatch.setenv("ENV", "production")
    r = TestClient(app).post("/api/projects/p-own/rooms/1/recalc?fixture=objektas_a", json={})
    assert r.status_code == 401, r.text


# ── -11 ──────────────────────────────────────────────────────────────────────

def test_telemetry_session_needs_login_and_owned_project(env, tmp_path, monkeypatch):
    monkeypatch.setenv("TELEMETRY_LOG_PATH", str(tmp_path / "tele.jsonl"))
    body = {"project_id": "p-own", "started_at": 1, "finished_at": 2, "total_seconds": 1,
            "gestures": []}
    assert TestClient(app).post("/api/telemetry/session", json=body).status_code == 401
    c = _login("owner@x.lt")
    assert c.post("/api/telemetry/session", json={**body, "project_id": "p-other"}).status_code == 404
    assert c.post("/api/telemetry/session",
                  json={**body, "gestures": [{}] * 2001}).status_code == 422
    assert c.post("/api/telemetry/session", json=body).status_code == 200
    assert (tmp_path / "tele.jsonl").read_text(encoding="utf-8").count("p-own") == 1


def test_telemetry_log_follows_data_dir(monkeypatch, tmp_path):
    from telemetry import session_log
    monkeypatch.delenv("TELEMETRY_LOG_PATH", raising=False)
    monkeypatch.setenv("SAMATU_DATA_DIR", str(tmp_path))
    assert session_log._log_path() == tmp_path / "telemetry" / "session_log.jsonl"
